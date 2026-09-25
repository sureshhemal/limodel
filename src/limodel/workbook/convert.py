"""Builds the four CSV files from the workbook, copying every value exactly as recorded.

Anything that looks wrong is reported, never corrected: the files stay a faithful copy of
what was written down. The layout is described in docs/data-format.md.
"""

import csv
import re
from dataclasses import dataclass

import openpyxl

from limodel.workbook.cells import MEASURED, POWER, TANKS, minute_rows, rinse_rows, run_sheets

RUN_FIELDS = ["run_code", "series", "run_date", "cem", "aem", "li_feed_mg_l", "mg_feed_mg_l",
              "licl_g_l", "mgcl2_6h2o_g_l", "conc_start_salt", "conc_start_g_l",
              "rinse_salt", "rinse_wt_pct", "set_current_ma", "flow_l_min", "area_cm2",
              "cell_pairs", "tare_conc_g", "tare_dil_g", "tare_rinse_g", "include", "notes"]

STACK_FIELDS = ["run_code", "minute"] + list(POWER) + ["notes"]

TANK_FIELDS = ["run_code", "minute", "tank", "conductivity_us_cm", "ph", "temp_c",
               "weight_g", "notes"]

SAMPLE_FIELDS = ["sample_id", "run_code", "minute", "tank", "li_mg_l", "mg_mg_l",
                 "na_mg_l", "k_mg_l", "ca_mg_l", "toc_ppm", "notes"]

MEASURED = ("li", "na", "mg", "k", "ca", "toc")


def plain(number):
    """Write 2.0 as 2, so values look as they do in the sheet."""
    return int(number) if float(number).is_integer() else number


def blank(value):
    """Excel cell to CSV value: None and empty strings become blank."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return value


def parse_conc_start(text):
    """'DI Water +0.5 g/l LiCl' -> ('LiCl', 0.5), keeping the salt spelled as recorded."""
    if not text:
        return "", ""
    match = re.search(r"([\d.]+)\s*g/?l\s*([A-Za-z0-9]+)", str(text), re.I)
    if not match:
        return "", ""
    amount, salt = match.groups()
    return salt, plain(float(amount))


def read_run(ws, code):
    date = ws.cell(1, 2).value
    salt, amount = parse_conc_start(ws.cell(3, 9).value)
    rinse_pct = re.search(r"[\d.]+", str(ws.cell(5, 2).value or ""))
    return {
        "run_code": code,
        "series": "new_set_mg",
        "run_date": date.date().isoformat() if hasattr(date, "date") else blank(date),
        "cem": blank(ws.cell(1, 17).value),
        "aem": blank(ws.cell(2, 17).value),
        "li_feed_mg_l": blank(ws.cell(1, 5).value),
        "mg_feed_mg_l": blank(ws.cell(2, 5).value),
        "licl_g_l": blank(ws.cell(3, 2).value),
        "mgcl2_6h2o_g_l": blank(ws.cell(2, 2).value),
        "conc_start_salt": salt,
        "conc_start_g_l": amount,
        "rinse_salt": blank(ws.cell(5, 1).value),
        "rinse_wt_pct": plain(float(rinse_pct.group())) if rinse_pct else "",
        "set_current_ma": "",
        "flow_l_min": blank(ws.cell(1, 9).value),
        "area_cm2": blank(ws.cell(2, 9).value),
        "cell_pairs": blank(ws.cell(1, 13).value),
        "tare_conc_g": blank(ws.cell(2, 25).value),
        "tare_dil_g": blank(ws.cell(3, 25).value),
        "tare_rinse_g": blank(ws.cell(4, 25).value),
        "include": "true",
        "notes": "",
    }


def read_stack(ws, row, code, minute):
    reading = {"run_code": code, "minute": minute, "notes": ""}
    for field, column in POWER.items():
        reading[field] = blank(ws.cell(row, column).value)
    return reading if any(reading[f] != "" for f in POWER) else None


def read_tank(ws, row, code, minute, tank, cols):
    reading = {
        "run_code": code,
        "minute": minute,
        "tank": tank,
        "conductivity_us_cm": blank(ws.cell(row, cols["conductivity"]).value),
        "ph": blank(ws.cell(row, cols["ph"]).value),
        "temp_c": blank(ws.cell(row, cols["temp"]).value),
        "weight_g": blank(ws.cell(row, cols["weight"]).value),
        "notes": "",
    }
    measured = [f"{m}={blank(ws.cell(row, cols[m]).value)}"
                for m in MEASURED if blank(ws.cell(row, cols[m]).value) != ""]
    if not blank(ws.cell(row, cols["id"]).value) and measured:
        # lab values with nothing to attach them to; kept here so nothing is lost
        reading["notes"] = "no sample ID for: " + ", ".join(measured)
    has_value = any(reading[f] != "" for f in
                    ("conductivity_us_cm", "ph", "temp_c", "weight_g", "notes"))
    return reading if has_value else None


def read_sample(ws, row, code, minute, tank, cols):
    sample_id = blank(ws.cell(row, cols["id"]).value)
    if not sample_id:
        return None
    return {
        "sample_id": sample_id,
        "run_code": code,
        "minute": minute,
        "tank": tank,
        "li_mg_l": blank(ws.cell(row, cols["li"]).value),
        "mg_mg_l": blank(ws.cell(row, cols["mg"]).value),
        "na_mg_l": blank(ws.cell(row, cols["na"]).value),
        "k_mg_l": blank(ws.cell(row, cols["k"]).value),
        "ca_mg_l": blank(ws.cell(row, cols["ca"]).value),
        "toc_ppm": blank(ws.cell(row, cols["toc"]).value),
        "notes": "",
    }


def convert_sheet(ws, code):
    """One sheet -> (run, stack readings, tank readings, samples, warnings)."""
    rinse = rinse_rows(ws)
    stack, tanks, samples, rinse_samples, warnings = [], [], [], [], []

    for minute, r in minute_rows(ws).items():
        reading = read_stack(ws, r, code, minute)
        if reading:
            stack.append(reading)

        rows = {"conc": r, "dil": r}
        if minute in rinse:
            rows["rinse"] = rinse[minute]

        for tank, tank_row in rows.items():
            tank_reading = read_tank(ws, tank_row, code, minute, tank, TANKS[tank])
            if tank_reading:
                tanks.append(tank_reading)
                if tank_reading["notes"]:
                    warnings.append(f"{code}: minute {minute} {tank} has "
                                    f"{tank_reading['notes']}")
            sample = read_sample(ws, tank_row, code, minute, tank, TANKS[tank])
            if sample:
                (rinse_samples if tank == "rinse" else samples).append(sample)

    samples.extend(rinse_samples)

    for sample in samples:
        if not sample["sample_id"].upper().startswith(code.upper()):
            warnings.append(f"{code}: sample ID {sample['sample_id']!r} does not start "
                            f"with {code}")
        if sample["li_mg_l"] == "" and sample["mg_mg_l"] == "":
            warnings.append(f"{code}: sample {sample['sample_id']} has no lab results yet")

    return read_run(ws, code), stack, tanks, samples, warnings


def write_csv(path, fields, rows):
    with open(path, "w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


@dataclass(frozen=True)
class Conversion:
    """What was written, and everything in the sheets that looks wrong."""
    runs: list[str]
    files: dict[str, int]
    warnings: list[str]


def convert(workbook_path, data_dir):
    """Rewrite the four CSV files from the workbook. Returns what was written.

    Existing files are replaced, so a value corrected in a CSV but not in the workbook is
    lost. Correct the workbook, not the CSV.
    """
    book = openpyxl.load_workbook(workbook_path, data_only=True)
    runs, stack, tanks, samples, warnings = [], [], [], [], []

    for code, ws in run_sheets(book):
        run, run_stack, run_tanks, run_samples, run_warnings = convert_sheet(ws, code)
        runs.append(run)
        stack.extend(run_stack)
        tanks.extend(run_tanks)
        samples.extend(run_samples)
        warnings.extend(run_warnings)

    tables = {"runs.csv": (RUN_FIELDS, runs),
              "stack_readings.csv": (STACK_FIELDS, stack),
              "tank_readings.csv": (TANK_FIELDS, tanks),
              "samples.csv": (SAMPLE_FIELDS, samples)}
    for name, (fields, rows) in tables.items():
        write_csv(data_dir / name, fields, rows)

    return Conversion(runs=[run["run_code"] for run in runs],
                      files={name: len(rows) for name, (_, rows) in tables.items()},
                      warnings=warnings)
