"""The convert command: rebuilds the data files from the workbook, changing nothing."""

import csv

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from limodel import config
from limodel.cli.main import app

runner = CliRunner()


@pytest.fixture
def convert(monkeypatch, workbook, tmp_path):
    """Runs the command against a folder with nothing in it yet."""
    out = tmp_path / "fresh"
    out.mkdir()

    def run(*arguments, into=out, input=None):
        monkeypatch.setattr(config, "WORKBOOK", workbook)
        monkeypatch.setattr(config, "DATA_DIR", into)
        return runner.invoke(app, ["convert", *arguments], input=input)
    return run


def rows(path):
    return list(csv.DictReader(open(path)))


def test_convert_writes_the_four_files(convert, tmp_path):
    result = convert()
    out = tmp_path / "fresh"

    assert result.exit_code == 0
    assert "converted 2 runs: AA BB" in result.output
    assert [row["run_code"] for row in rows(out / "runs.csv")] == ["AA", "BB"]
    assert [row["minute"] for row in rows(out / "stack_readings.csv")] == ["0", "5", "0"]
    assert [row["tank"] for row in rows(out / "tank_readings.csv")] == \
           ["conc", "dil", "rinse", "conc"]
    assert [row["sample_id"] for row in rows(out / "samples.csv")] == ["AA1", "AA2", "BB1"]


def test_convert_copies_a_value_exactly_as_recorded(convert, tmp_path):
    convert()

    conc = rows(tmp_path / "fresh" / "tank_readings.csv")[0]

    assert conc["weight_g"] == "2119.7"
    assert conc["temp_c"] == "21.4"
    assert conc["notes"] == ""


def test_convert_reports_a_bad_sample_id_instead_of_fixing_it(convert, workbook, tmp_path):
    book = load_workbook(workbook)
    book["New_set_Mg2+(07_10_26)_AA"].cell(9, 13).value = "XX1"
    book.save(workbook)

    result = convert()

    assert "does not start with AA" in result.output
    assert "Nothing was corrected" in result.output
    assert rows(tmp_path / "fresh" / "samples.csv")[0]["sample_id"] == "XX1"


def test_convert_asks_before_replacing_files_that_are_there(convert, data_dir):
    before = (data_dir / "runs.csv").read_text()

    result = convert(into=data_dir, input="n\n")

    assert result.exit_code == 1
    assert "This replaces 4 files" in result.output
    assert (data_dir / "runs.csv").read_text() == before


def test_convert_does_not_ask_when_forced(convert, data_dir):
    result = convert("--force", into=data_dir)

    assert result.exit_code == 0
    assert "This replaces" not in result.output
    assert [row["run_code"] for row in rows(data_dir / "runs.csv")] == ["AA", "BB"]
