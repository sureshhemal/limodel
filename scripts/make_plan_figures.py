"""Make the figures in docs/plan-for-building-the-model.md from the lab's real GO-10 runs.

Every number is read from the lab's workbook (data/ED_data_from M-BF(Updated).xlsx): each run's
whole-run lithium recovery, its planned Mg/Li ratio, its current density and its highest stack
voltage. The runs are put through Steps 2 to 8 of the plan for lithium recovery. The figures are
written into docs/figures/plan (or the folder given as the first argument), and the numbers the
plan quotes are printed as JSON.

XGBoost (Step 7) is not a project dependency yet, so run it with:

    uv run --with xgboost python scripts/make_plan_figures.py
"""

import json
import math
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from scipy import stats
from xgboost import XGBRegressor

ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = ROOT / "data" / "ED_data_from M-BF(Updated).xlsx"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "docs" / "figures" / "plan"
OUT.mkdir(parents=True, exist_ok=True)

# --- palette (reference instance, light surface) -------------------------------------------
SURFACE, INK, INK_2, GRID = "#fcfcfb", "#0b0b0b", "#52514e", "#e4e3df"
BLUE, ORANGE = "#2a78d6", "#eb6834"
BAND = "#cde2fb"

plt.rcParams.update({
    "figure.facecolor": SURFACE, "axes.facecolor": SURFACE, "savefig.facecolor": SURFACE,
    "axes.edgecolor": GRID, "axes.labelcolor": INK_2, "xtick.color": INK_2, "ytick.color": INK_2,
    "text.color": INK, "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.8,
    "axes.spines.top": False, "axes.spines.right": False, "font.size": 10,
    "axes.titlesize": 11, "axes.titleweight": "bold", "axes.titlecolor": INK,
    "lines.linewidth": 2, "lines.markersize": 8, "legend.frameon": False,
})


MASS_PER_MOLAR = 24.305 / 6.94  # a molar Mg/Li ratio times this is the mass ratio


def mass_ratio_axis(ax):
    ax.set_xscale("log")
    ax.set_xticks([5, 10, 20, 40, 60])
    ax.set_xticklabels(["5", "10", "20", "40", "60"])
    ax.minorticks_off()


def save(fig, name):
    fig.savefig(OUT / name, dpi=160, bbox_inches="tight")
    plt.close(fig)


# --- the runs, from the workbook ---------------------------------------------------------------
# The GO-10 runs at 400 mg/L lithium that passed the standard charge, and the cell on each run's
# sheet that holds its whole-run recovery (last concentrate sample against the first).
WHOLE_RUN_RECOVERY = {"AJ": "AZ44", "AH": "AZ33", "AI": "AZ33", "AE": "AZ28", "BD": "AZ28",
                      "BE": "AZ28", "AF": "AZ28", "AG": "AZ28", "AT": "AZ28", "AW": "AZ28",
                      "AN": "AZ17", "AP": "AZ17"}
STANDARD_CONCENTRATE = "DI Water +0.5 g/l LiCl"
REPEAT_GROUPS = [["AH", "AI"], ["AE", "BD", "BE"]]

book = openpyxl.load_workbook(WORKBOOK, data_only=True)
runs = []
for code, cell in WHOLE_RUN_RECOVERY.items():
    sheet = book[next(n for n in book.sheetnames if n.startswith("New_set") and n.endswith("_" + code))]
    last_row = next(r for r in range(9, sheet.max_row + 1) if sheet.cell(r, 1).value == "Rinse")
    hourly_density = [sheet.cell(r, 8).value for r in range(9, last_row)
                      if isinstance(sheet.cell(r, 8).value, (int, float))]
    voltages = [sheet.cell(r, 5).value for r in range(9, last_row)
                if isinstance(sheet.cell(r, 5).value, (int, float))]
    runs.append({"run": code, "i": round(sum(hourly_density) / len(hourly_density), 2),
                 "R": sheet["E4"].value, "mass_ratio": sheet["E3"].value,
                 "recovery": round(sheet[cell].value, 2), "U_max": max(voltages),
                 "standard": sheet["I3"].value == STANDARD_CONCENTRATE})
numbers = {"runs": runs}
n = len(runs)
y_rec = np.array([r["recovery"] for r in runs])


# --- Step 3: pure error --------------------------------------------------------------------------
def pure_error(values):
    by_run = dict(zip([r["run"] for r in runs], values))
    ss, nu = 0.0, 0
    for group in REPEAT_GROUPS:
        ys = [by_run[c] for c in group]
        ss += sum((v - np.mean(ys)) ** 2 for v in ys)
        nu += len(ys) - 1
    return ss, nu


ss_pe, nu_pe = pure_error(y_rec)
s_pe = math.sqrt(ss_pe / nu_pe)
in_groups = [r["recovery"] for r in runs if any(r["run"] in g for g in REPEAT_GROUPS)]
numbers["pure_error"] = {"ss": ss_pe, "nu": nu_pe, "s": s_pe, "mean": float(np.mean(in_groups)),
                         "relative": s_pe / np.mean(in_groups),
                         "lo": s_pe * math.sqrt(nu_pe / stats.chi2.ppf(0.975, nu_pe)),
                         "hi": s_pe * math.sqrt(nu_pe / stats.chi2.ppf(0.025, nu_pe)),
                         "groups": {"+".join(g): [next(r["recovery"] for r in runs if r["run"] == c) for c in g]
                                    for g in REPEAT_GROUPS}}


# --- Step 4: coding ------------------------------------------------------------------------------
def coder(values):
    lo, hi = min(values), max(values)
    return (hi + lo) / 2, (hi - lo) / 2


i_all = np.array([r["i"] for r in runs])
ln_r_all = np.log([r["R"] for r in runs])
CODING = {"i": coder(i_all), "ln i": coder(np.log(i_all)), "ln R": coder(ln_r_all)}
xi = (i_all - CODING["i"][0]) / CODING["i"][1]
numbers["coding"] = {k: list(v) for k, v in CODING.items()}
numbers["coding"]["corr_raw"] = float(np.corrcoef(i_all, i_all ** 2)[0, 1])
numbers["coding"]["corr_coded"] = float(np.corrcoef(xi, xi ** 2)[0, 1])


def coded(i, r, scale):
    v1 = math.log(i) if scale == "ln i" else i
    return (v1 - CODING[scale][0]) / CODING[scale][1], (math.log(r) - CODING["ln R"][0]) / CODING["ln R"][1]


# --- Step 5: the ladder, ratio first -------------------------------------------------------------
TERMS = {"M0": [], "M1": ["x2"], "M2": ["x2", "x1"], "M3": ["x2", "x1", "x1^2"],
         "M4": ["x2", "x1", "x1^2", "x2^2"]}


def matrix(rows, terms, scale):
    out = []
    for row in rows:
        x1, x2 = coded(row["i"], row["R"], scale)
        cols = {"x1": x1, "x2": x2, "x1^2": x1 ** 2, "x2^2": x2 ** 2}
        out.append([1.0] + [cols[t] for t in terms])
    return np.array(out)


def fit(rows, y, terms, scale, log_output):
    X = matrix(rows, terms, scale)
    k, p = X.shape
    XtX_inv = np.linalg.inv(X.T @ X)
    b = XtX_inv @ X.T @ y
    e = y - X @ b
    ss_res, ss_tot = float(e @ e), float(((y - y.mean()) ** 2).sum())
    s2 = ss_res / (k - p)
    h = np.einsum("ij,jk,ik->i", X, XtX_inv, X)
    loo_pred = y - e / (1 - h)
    measured = np.exp(y) if log_output else y
    loo_pred_rec = np.exp(loo_pred) if log_output else loo_pred
    miss = measured - loo_pred_rec
    se = np.sqrt(np.diag(XtX_inv) * s2)
    out = {"terms": terms, "scale": scale, "log": log_output, "n": k, "p": p, "b": b, "se": se,
           "XtX_inv": XtX_inv, "s2": s2, "ss_res": ss_res, "e": e, "h": h, "y": y,
           "loo_pred_rec": loo_pred_rec, "rmse_loo": float(np.sqrt((miss ** 2).mean())),
           "mae_loo": float(np.abs(miss).mean()), "q2": 1 - float((miss ** 2).sum()) / float(((measured - measured.mean()) ** 2).sum()),
           "r2": 1 - ss_res / ss_tot, "r2_adj": 1 - (ss_res / ss_tot) * (k - 1) / (k - p),
           "t_p": [float(2 * stats.t.sf(abs(bj / sj), k - p)) for bj, sj in zip(b, se)],
           "cook": (e ** 2 / (p * s2)) * h / (1 - h) ** 2}
    if p > 1:
        F = ((ss_tot - ss_res) / (p - 1)) / s2
        out["F"], out["F_p"] = F, float(stats.f.sf(F, p - 1, k - p))
    ss_pe_y, nu_pe_y = pure_error(y)
    df_lof = k - p - nu_pe_y
    f_lof = ((ss_res - ss_pe_y) / df_lof) / (ss_pe_y / nu_pe_y)
    out["lof_F"], out["lof_p"] = f_lof, float(stats.f.sf(f_lof, df_lof, nu_pe_y))
    return out


def ladder(y, scale, log_output):
    fits, prev = {}, None
    for label, terms in TERMS.items():
        f = fit(runs, y, terms, scale, log_output)
        if prev is not None:
            Fp = ((prev["ss_res"] - f["ss_res"]) / (f["p"] - prev["p"])) / (f["ss_res"] / (f["n"] - f["p"]))
            f["partial_F"], f["partial_p"] = Fp, float(stats.f.sf(Fp, f["p"] - prev["p"], f["n"] - f["p"]))
        fits[label], prev = f, f
    return fits


def choose(fits, s_pe_rec):
    ok = [k for k, f in fits.items()
          if all(pv < 0.05 for pv in f["t_p"][1:]) and f["lof_p"] >= 0.05 and f["p"] <= f["n"] - 2]
    best = min(ok, key=lambda k: fits[k]["rmse_loo"])
    close = [k for k in ok if fits[k]["rmse_loo"] - fits[best]["rmse_loo"] <= s_pe_rec]
    return min(close, key=lambda k: fits[k]["p"])


def summary(fits):
    return {k: {"p": f["p"], "R2": f["r2"], "R2_adj": f["r2_adj"], "F": f.get("F"), "F_p": f.get("F_p"),
                "partial_F": f.get("partial_F"), "partial_p": f.get("partial_p"), "rmse_loo": f["rmse_loo"],
                "mae_loo": f["mae_loo"], "lof_F": f["lof_F"], "lof_p": f["lof_p"], "t_p": f["t_p"],
                "b": [round(float(v), 3) for v in f["b"]]} for k, f in fits.items()}


plain = {scale: ladder(y_rec, scale, False) for scale in ("i", "ln i")}
logged = {scale: ladder(np.log(y_rec), scale, True) for scale in ("i", "ln i")}
numbers["ladder_plain"] = {scale: summary(f) for scale, f in plain.items()}
numbers["ladder_log"] = {scale: summary(f) for scale, f in logged.items()}
numbers["choice_plain"] = {scale: choose(f, s_pe) for scale, f in plain.items()}
numbers["choice_log"] = {scale: choose(f, s_pe) for scale, f in logged.items()}
first = plain["ln i"][numbers["choice_plain"]["ln i"]]
final = logged["ln i"][numbers["choice_log"]["ln i"]]
numbers["s_pe_log"] = math.sqrt(pure_error(np.log(y_rec))[0] / nu_pe)


def predict(f, i, r):
    x0 = matrix([{"i": i, "R": r}], f["terms"], f["scale"])[0]
    y0 = float(x0 @ f["b"])
    h0 = float(x0 @ f["XtX_inv"] @ x0)
    half = stats.t.ppf(0.975, f["n"] - f["p"]) * math.sqrt(f["s2"] * (1 + h0))
    if f["log"]:
        return math.exp(y0), math.exp(y0 - half), math.exp(y0 + half), h0
    return y0, y0 - half, y0 + half, h0


numbers["first_at_ratio_60"] = predict(first, 12, max(r["R"] for r in runs))
numbers["final"] = {"b": [float(v) for v in final["b"]], "se": [float(v) for v in final["se"]],
                    "t_p": final["t_p"], "s": math.sqrt(final["s2"]),
                    "exponent_of_R": float(final["b"][1] / CODING["ln R"][1]),
                    "constant": float(math.exp(final["b"][0] - final["b"][1] * CODING["ln R"][0] / CODING["ln R"][1]))}

# --- Step 6: checks of the final model -----------------------------------------------------------
std_res = final["e"] / np.sqrt(final["s2"] * (1 - final["h"]))
ranks = np.empty(n, int)
ranks[np.argsort(std_res)] = np.arange(1, n + 1)
z_expected = stats.norm.ppf((ranks - 0.5) / n)
flagged = [r["run"] for r, h, c in zip(runs, final["h"], final["cook"])
           if h > 2 * final["p"] / n or c > 4 / n]
refits = {}
for drop in flagged:
    kept = [r for r in runs if r["run"] != drop]
    X = matrix(kept, final["terms"], "ln i")
    b_without = np.linalg.lstsq(X, np.log([r["recovery"] for r in kept]), rcond=None)[0]
    refits[drop] = {"b_without": [round(float(v), 3) for v in b_without],
                    "shift_in_se": [round(float((b - bw) / se), 2)
                                    for b, bw, se in zip(final["b"], b_without, final["se"])]}
first_std_res = first["e"] / np.sqrt(first["s2"] * (1 - first["h"]))
numbers["checks"] = {"leverage_limit": 2 * final["p"] / n, "cook_limit": 4 / n, "flagged": flagged,
                     "refits": refits, "shapiro_p": float(stats.shapiro(final["e"]).pvalue),
                     "rows": [{"run": r["run"], "h": round(float(h), 3), "r": round(float(sr), 2),
                               "cook": round(float(c), 3), "loo_pred": round(float(lp), 1),
                               "first_r": round(float(fr), 2)}
                              for r, h, sr, c, lp, fr in zip(runs, final["h"], std_res, final["cook"],
                                                             final["loo_pred_rec"], first_std_res)]}

# --- Step 7: boosted trees, leave-one-out ------------------------------------------------------
X_raw = np.array([[r["i"], r["R"]] for r in runs])


def trees():
    return XGBRegressor(n_estimators=100, max_depth=2, random_state=0)


tree_loo = np.array([float(trees().fit(np.delete(X_raw, k, 0), np.delete(y_rec, k)).predict(X_raw[[k]])[0])
                     for k in range(n)])
tree_miss = y_rec - tree_loo
numbers["trees"] = {"rmse_loo": float(np.sqrt((tree_miss ** 2).mean())), "mae_loo": float(np.abs(tree_miss).mean()),
                    "q2": 1 - float((tree_miss ** 2).sum()) / float(((y_rec - y_rec.mean()) ** 2).sum()),
                    "loo": {r["run"]: round(float(t), 1) for r, t in zip(runs, tree_loo)},
                    "poly": {"rmse_loo": final["rmse_loo"], "mae_loo": final["mae_loo"], "q2": final["q2"]}}
tree_all = trees().fit(X_raw, y_rec)
ratio_grid = np.exp(np.linspace(math.log(1.2), math.log(20), 300))

# --- Step 8: the best setup ------------------------------------------------------------------------
r_lo, r_hi = min(r["R"] for r in runs), max(r["R"] for r in runs)
i_lo, i_hi = float(i_all.min()), float(i_all.max())
grid = [(i, r) for i in np.arange(round(i_lo, 1), i_hi + 1e-9, 0.1)
        for r in np.exp(np.linspace(math.log(r_lo), math.log(r_hi), 60))]
best_i, best_r = max(grid, key=lambda g: predict(final, *g)[0])
best = predict(final, best_i, best_r)
window_ratios = [r for r in np.exp(np.linspace(math.log(r_lo), math.log(r_hi), 400))
                 if predict(final, 12, r)[0] >= best[0] - s_pe]
charge_per_area = 48.5  # mA·h/cm², the same in every run
voltages = {}
for r in runs:
    key = (round(r["i"] / 3) * 3 if r["i"] < 4 else round(r["i"] / 6) * 6, r["standard"])
    voltages.setdefault(key, []).append(r["U_max"])
numbers["best"] = {"i_searched": [i_lo, i_hi], "R_searched": [r_lo, r_hi], "R": best_r, "mass_ratio": best_r * MASS_PER_MOLAR,
                   "recovery": best[0], "interval": [best[1], best[2]], "h0": best[3],
                   "window_upper_R": max(window_ratios), "window_upper_mass": max(window_ratios) * MASS_PER_MOLAR,
                   "flat_in_i": len(final["terms"]) == 1,
                   "voltages": {f"{k[0]} mA/cm2, {'standard' if k[1] else 'other'} concentrate": v
                                for k, v in sorted(voltages.items())},
                   "run_length_h": {lvl: charge_per_area / lvl for lvl in (3, 6, 12, 24)},
                   "interval_parts": {"t": float(stats.t.ppf(0.975, final["n"] - final["p"])),
                                      "s": math.sqrt(final["s2"])}}

# --- Step 9: predictions written down for the next runs --------------------------------------
NEXT_RUNS = [("12 mA/cm², ratio 5", 12, 5), ("3 mA/cm², ratio 5", 3, 5), ("6 mA/cm², ratio 20", 6, 20)]
numbers["next_runs"] = [{"setting": label, "prediction": [round(v, 1) for v in predict(final, i, mass / MASS_PER_MOLAR)[:3]]}
                        for label, i, mass in NEXT_RUNS]

# ================================ figures ==================================================
colour = [BLUE if r["standard"] else ORANGE for r in runs]

# Step 2a: the design
fig, ax = plt.subplots(figsize=(5.8, 3.9))
cells = {}
for r in runs:
    key = (round(r["i"] / 3) * 3 if r["i"] < 4 else round(r["i"] / 6) * 6, r["mass_ratio"], r["standard"])
    cells.setdefault(key, []).append(r["run"])
for (lvl, mass, std), codes in cells.items():
    ax.scatter(lvl, mass, s=60 + 50 * (len(codes) - 1), color=BLUE if std else ORANGE, zorder=3,
               edgecolor=SURFACE, linewidth=2)
    ax.annotate(", ".join(codes), (lvl, mass), textcoords="offset points",
                xytext=(9, -12 if not std else 6), color=INK_2, fontsize=8)
ax.set_xscale("log", base=2); ax.set_yscale("log", base=2)
ax.set_xticks([3, 6, 12, 24]); ax.set_xticklabels(["3", "6", "12", "24"])
ax.set_yticks([5, 10, 20, 40, 60]); ax.set_yticklabels(["5", "10", "20", "40", "60"])
ax.set_xlim(2.4, 36)
ax.set_xlabel("current density, mA/cm²"); ax.set_ylabel("Mg/Li ratio (mass)")
ax.scatter([], [], color=BLUE, label="standard concentrate"); ax.scatter([], [], color=ORANGE, label="different concentrate")
ax.legend(loc="upper right", fontsize=8)
ax.set_title("The GO-10 runs: a cross around the base run")
save(fig, "step2-design.png")

# Step 2b: shapes of recovery
arm_i = [r for r in runs if r["mass_ratio"] == 5]
arm_r = [r for r in runs if abs(r["i"] - 12) < 1]
fig, axes = plt.subplots(1, 3, figsize=(11, 3.4), sharey=True)
axes[0].scatter([r["i"] for r in arm_i], [r["recovery"] for r in arm_i],
                color=[BLUE if r["standard"] else ORANGE for r in arm_i], zorder=3)
axes[0].set_xlabel("current density, mA/cm² (ratio 5)"); axes[0].set_ylabel("lithium recovery, %")
axes[0].set_title("Recovery against current density")
axes[1].scatter([r["mass_ratio"] for r in arm_r], [r["recovery"] for r in arm_r], color=BLUE, zorder=3)
axes[1].set_xlabel("Mg/Li ratio, mass (12 mA/cm²)"); axes[1].set_title("Recovery against ratio")
axes[2].scatter([math.log(r["R"]) for r in arm_r], [r["recovery"] for r in arm_r], color=BLUE, zorder=3)
axes[2].set_xlabel("ln R, molar (12 mA/cm²)"); axes[2].set_title("Recovery against ln R")
fig.suptitle("Lithium recovery against each input (GO-10; orange: different concentrate)", fontweight="bold", y=1.04)
save(fig, "step2-shapes.png")

# Step 3: repeat groups
fig, ax = plt.subplots(figsize=(5.6, 3.4))
for k, group in enumerate(REPEAT_GROUPS):
    members = sorted([r for r in runs if r["run"] in group], key=lambda r: r["recovery"])
    ys = [m["recovery"] for m in members]
    ax.scatter([k] * len(ys), ys, color=BLUE, zorder=3)
    for n_m, m in enumerate(members):
        side = 1 if n_m % 2 == 0 else -1
        ax.annotate(m["run"], (k, m["recovery"]), textcoords="offset points", xytext=(10 * side, -3),
                    ha="left" if side > 0 else "right", color=INK_2)
    ax.hlines(np.mean(ys), k - 0.18, k + 0.18, color=INK, linewidth=2)
ax.set_xticks([0, 1]); ax.set_xticklabels(["6 mA/cm², ratio 5", "12 mA/cm², ratio 5"])
ax.set_xlim(-0.6, 1.6)
ax.set_ylabel("lithium recovery, %"); ax.grid(axis="x", visible=False)
ax.set_title("GO-10 repeat runs and their group averages")
save(fig, "step3-repeats.png")

# Step 4: overlap before and after coding
fig, axes = plt.subplots(1, 2, figsize=(8.4, 3.2))
axes[0].scatter(i_all, i_all ** 2, color=BLUE, zorder=3)
axes[0].set_xlabel("i"); axes[0].set_ylabel("i²")
axes[0].set_title(f"Raw: correlation {numbers['coding']['corr_raw']:.2f}")
axes[1].scatter(xi, xi ** 2, color=BLUE, zorder=3)
axes[1].set_xlabel("coded x"); axes[1].set_ylabel("x²")
axes[1].set_title(f"Coded: correlation {numbers['coding']['corr_coded']:.2f}")
fig.suptitle("Current density and its square, before and after coding (GO-10 runs)", fontweight="bold", y=1.04)
save(fig, "step4-coding.png")

# Step 5: leave-one-out error up the ladder, both output scales
fig, ax = plt.subplots(figsize=(6.4, 3.5))
labels = list(TERMS)
xs = np.arange(len(labels))
for offset, fits, name, col in ((-0.2, plain["ln i"], "recovery as measured", BLUE),
                                (0.2, logged["ln i"], "log of recovery", ORANGE)):
    vals = [fits[k]["rmse_loo"] for k in labels]
    bars = ax.bar(xs + offset, vals, width=0.38, color=col, label=name)
    for bb, v in zip(bars, vals):
        ax.annotate(f"{v:.1f}", (bb.get_x() + bb.get_width() / 2, v), textcoords="offset points",
                    xytext=(0, 3), ha="center", color=INK_2, fontsize=8)
ax.axhline(s_pe, color=INK, linestyle="--", linewidth=1.5)
ax.annotate(f"pure error\n{s_pe:.1f}", xy=(1.0, s_pe), xycoords=("axes fraction", "data"),
            textcoords="offset points", xytext=(6, 0), ha="left", va="center", color=INK)
ax.set_xticks(xs); ax.set_xticklabels(labels)
ax.set_ylabel("leave-one-out RMSE, points of recovery"); ax.grid(axis="x", visible=False)
ax.legend(loc="upper right", fontsize=8)
ax.set_title("Prediction error up the ladder (GO-10 runs)")
save(fig, "step5-ladder.png")

# Step 6a: why the log scale
fig, axes = plt.subplots(1, 2, figsize=(10.5, 3.6))
ax = axes[0]
ax.axhline(0, color=INK_2, linewidth=1)
ax.scatter([r["mass_ratio"] for r in runs], first_std_res, color=colour, zorder=3)
mass_ratio_axis(ax)
ax.set_xlabel("Mg/Li ratio (mass)"); ax.set_ylabel("standardised residual")
ax.set_title("Recovery as measured: misses against ratio")
ax = axes[1]
ax.scatter([r["mass_ratio"] for r in arm_r], [r["recovery"] for r in arm_r], color=INK, zorder=4, s=36, label="runs at 12 mA/cm²")
ax.plot(ratio_grid * MASS_PER_MOLAR, [predict(first, 12, r)[0] for r in ratio_grid], color=BLUE, label="fitted to recovery as measured")
ax.plot(ratio_grid * MASS_PER_MOLAR, [predict(final, 12, r)[0] for r in ratio_grid], color=ORANGE, label="fitted to the log of recovery")
ax.axhline(0, color=INK_2, linewidth=1)
ax.axvspan(r_lo * MASS_PER_MOLAR, r_hi * MASS_PER_MOLAR, color=BAND, alpha=0.4, linewidth=0)
mass_ratio_axis(ax); ax.set_xlabel("Mg/Li ratio, mass (shaded: tested)"); ax.set_ylabel("lithium recovery, %")
ax.set_title("The two fits along the ratio")
ax.legend(loc="upper right", fontsize=8)
fig.suptitle("Why lithium recovery is modelled on a log scale (GO-10 runs)", fontweight="bold", y=1.03)
fig.tight_layout()
save(fig, "step6-log-scale.png")

# Step 6b: checks of the final model
fig, axes = plt.subplots(2, 2, figsize=(9.6, 7.4))
ax = axes[0, 0]
lim = [0, 50]
ax.fill_between(lim, [v - 2 * s_pe for v in lim], [v + 2 * s_pe for v in lim], color=BAND, linewidth=0)
ax.plot(lim, lim, color=INK_2, linewidth=1)
ax.scatter(final["loo_pred_rec"], y_rec, color=colour, zorder=3)
for r, lp, sr in zip(runs, final["loo_pred_rec"], std_res):
    if abs(sr) > 1:
        ax.annotate(r["run"], (lp, r["recovery"]), textcoords="offset points", xytext=(6, -3), fontsize=8, color=INK_2)
ax.set_xlim(lim); ax.set_ylim(lim)
ax.set_xlabel("predicted recovery, run left out (%)"); ax.set_ylabel("measured recovery (%)")
ax.set_title("Parity (band: ±2 pure errors)")
ax = axes[0, 1]
ax.axhspan(-2, 2, color=BAND, linewidth=0)
ax.axhline(0, color=INK_2, linewidth=1)
ax.scatter([r["i"] for r in runs], std_res, color=colour, zorder=3)
for r, sr in zip(runs, std_res):
    if abs(sr) > 1:
        ax.annotate(r["run"], (r["i"], sr), textcoords="offset points", xytext=(6, -3), fontsize=8, color=INK_2)
ax.set_ylim(-3, 3)
ax.set_xlabel("current density, mA/cm²"); ax.set_ylabel("standardised residual")
ax.set_title("Misses against current density")
ax = axes[1, 0]
ax.plot([-2, 2], [-2, 2], color=INK_2, linewidth=1)
ax.scatter(z_expected, std_res, color=colour, zorder=3)
ax.set_xlabel("expected if normal"); ax.set_ylabel("standardised residual"); ax.set_title("Normal Q–Q")
ax = axes[1, 1]
ax.bar([r["run"] for r in runs], final["cook"], color=colour, width=0.6)
ax.axhline(4 / n, color=INK, linestyle="--", linewidth=1.5)
ax.annotate("4/n", (-0.4, 4 / n), textcoords="offset points", xytext=(0, 4), ha="left", color=INK)
ax.set_ylabel("Cook's distance"); ax.set_title("Influence of each run")
ax.tick_params(axis="x", labelsize=8); ax.grid(axis="x", visible=False)
fig.suptitle("Checking the chosen model for lithium recovery (GO-10; orange: different concentrate)",
             fontweight="bold", y=1.0)
fig.tight_layout()
save(fig, "step6-checks.png")

# Step 7: polynomial against boosted trees
fig, axes = plt.subplots(1, 2, figsize=(11.0, 3.6), gridspec_kw={"width_ratios": [1, 1.5]})
ax = axes[0]
vals = [final["rmse_loo"], numbers["trees"]["rmse_loo"]]
bars = ax.bar(["polynomial", "boosted trees"], vals, color=BLUE, width=0.5)
for bb, v in zip(bars, vals):
    ax.annotate(f"{v:.1f}", (bb.get_x() + bb.get_width() / 2, v), textcoords="offset points",
                xytext=(0, 4), ha="center", color=INK_2)
ax.axhline(s_pe, color=INK, linestyle="--", linewidth=1.5)
ax.annotate(f"pure error\n{s_pe:.1f}", xy=(1.0, s_pe), xycoords=("axes fraction", "data"),
            textcoords="offset points", xytext=(6, 0), ha="left", va="center", color=INK)
ax.set_ylabel("leave-one-out RMSE, points of recovery"); ax.grid(axis="x", visible=False)
ax.set_title("Prediction error on unseen runs")
ax = axes[1]
ax.axvspan(r_lo * MASS_PER_MOLAR, r_hi * MASS_PER_MOLAR, color=BAND, linewidth=0, alpha=0.6)
ax.plot(ratio_grid * MASS_PER_MOLAR, [predict(final, 12, r)[0] for r in ratio_grid], color=BLUE, label="polynomial")
ax.plot(ratio_grid * MASS_PER_MOLAR, tree_all.predict(np.column_stack([np.full_like(ratio_grid, 12.0), ratio_grid])),
        color=ORANGE, label="boosted trees", drawstyle="steps-mid")
ax.scatter([r["mass_ratio"] for r in arm_r], [r["recovery"] for r in arm_r], color=INK, zorder=4, s=36, label="runs at 12 mA/cm²")
mass_ratio_axis(ax)
ax.set_xlabel("Mg/Li ratio, mass (shaded: tested range)"); ax.set_ylabel("predicted recovery at 12 mA/cm², %")
ax.set_title("Both models fitted to all runs")
ax.legend(loc="upper right")
fig.suptitle("Lithium recovery: polynomial against a flexible model (GO-10 runs)", fontweight="bold", y=1.03)
fig.tight_layout()
save(fig, "step7-comparison.png")

# Step 8: the best setup
fig, axes = plt.subplots(1, 2, figsize=(11.0, 3.8))
ax = axes[0]
pred = [predict(final, 12, r) for r in ratio_grid]
inside = (ratio_grid >= r_lo) & (ratio_grid <= r_hi)
ax.fill_between(ratio_grid[inside] * MASS_PER_MOLAR, [p[1] for p, k in zip(pred, inside) if k], [p[2] for p, k in zip(pred, inside) if k],
                color=BAND, linewidth=0, label="95 % prediction interval")
ax.plot(ratio_grid[inside] * MASS_PER_MOLAR, [p[0] for p, k in zip(pred, inside) if k], color=BLUE, label="predicted recovery")
ax.scatter([r["mass_ratio"] for r in runs], y_rec, color=colour, zorder=4, s=30, label="runs (orange: different concentrate)")
ax.scatter([best_r * MASS_PER_MOLAR], [best[0]], color=INK, zorder=5, s=70, marker="D")
ax.annotate(f"best: ratio 5, {best[0]:.0f} %", (best_r * MASS_PER_MOLAR, best[0]), textcoords="offset points", xytext=(10, 4), color=INK)
mass_ratio_axis(ax); ax.set_xlabel("Mg/Li ratio (mass)"); ax.set_ylabel("lithium recovery, %")
ax.set_title("Predicted recovery: the same at every current density")
ax.legend(loc="upper right", fontsize=8)
ax = axes[1]
ax.axhline(20, color=INK, linestyle="--", linewidth=1.5)
ax.annotate("20 V: the stack's limit", (2.6, 20), textcoords="offset points", xytext=(0, 4), color=INK)
ax.scatter([r["i"] for r in runs], [r["U_max"] for r in runs], color=colour, zorder=3)
for r in runs:
    if not r["standard"]:
        ax.annotate(r["run"], (r["i"], r["U_max"]), textcoords="offset points", xytext=(6, -3), fontsize=8, color=INK_2)
ax.set_xscale("log", base=2); ax.set_xticks([3, 6, 12, 24]); ax.set_xticklabels(["3", "6", "12", "24"])
ax.set_ylim(0, 22)
ax.set_xlabel("current density, mA/cm²"); ax.set_ylabel("highest stack voltage, V")
ax.set_title("Measured highest voltage of each run")
fig.suptitle("Finding the best setup (GO-10 runs)", fontweight="bold", y=1.03)
fig.tight_layout()
save(fig, "step8-best-setup.png")


def plain_numbers(o):
    if isinstance(o, dict):
        return {str(k): plain_numbers(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return [plain_numbers(v) for v in o]
    if isinstance(o, np.generic):
        return o.item()
    if isinstance(o, np.ndarray):
        return o.tolist()
    return o


print(json.dumps(plain_numbers(numbers), indent=1))
